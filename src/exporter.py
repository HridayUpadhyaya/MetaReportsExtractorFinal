from __future__ import annotations
from collections import defaultdict
from copy import copy
from datetime import date, datetime, timezone
from pathlib import Path
import re
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment


def _copy_style_row(ws, source_row: int, target_row: int, max_col: int):
    for c in range(1, max_col + 1):
        src = ws.cell(source_row, c)
        dst = ws.cell(target_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)


def _proactive(r: dict) -> int:
    return int(round(float(r.get("content_actioned_numeric") or 0) * float(r.get("proactive_rate") or 0)))


def _excel_date(value):
    return date.fromisoformat(value) if value else None


def _failed_report_to_text(item) -> str:
    """Convert failed-report records to an Excel-safe human-readable string."""
    if isinstance(item, dict):
        publication = item.get("publication_month") or item.get("edition_key") or "Unknown report"
        error = item.get("error") or "Unknown error"
        issues = item.get("audit", {}).get("validation", {}).get("report_issues", [])
        if "no_policy_rows" in issues:
            error = "No content-policy table is present in this report."
        elif any(issue.startswith("expected_") for issue in issues):
            counts = []
            for issue in issues:
                match = re.fullmatch(r"expected_(\d+)_policy_rows:(\w+):found_(\d+)", issue)
                if match:
                    expected, platform, found = match.groups()
                    counts.append(f"{platform}: prose states {expected} policy areas, table contains {found}")
            error = "; ".join(counts) or "Policy counts did not match the source."
        elif "Review file:" in error:
            error = "Failed strict validation; see the saved audit."
        return f"{publication}: {error}"
    if item is None:
        return "Unknown failure"
    return str(item)


def build_workbook(template: Path, output: Path, rows: list[dict], failed_reports: list[str] | None = None) -> None:
    failed_reports = failed_reports or []
    porder = {"Facebook": 0, "Instagram": 1, "Threads": 2}
    rows = sorted(rows, key=lambda r: (r.get("period_end") or "", porder.get(r.get("platform"), 99), r.get("policy_order", 999), r.get("policy_category") or ""))
    wb = load_workbook(template)

    md = wb["Master Data"]
    # The template's old note can sit inside the expanded data range. Unmerge
    # it before adding rows or Excel will discard that row's other cells.
    for merged_range in list(md.merged_cells.ranges):
        md.unmerge_cells(str(merged_range))
    style_row = 2
    styles = [copy(md.cell(style_row, c)._style) for c in range(1, 12)]
    formats = [md.cell(style_row, c).number_format for c in range(1, 12)]
    if md.max_row > 1:
        md.delete_rows(2, md.max_row - 1)
    for i, r in enumerate(rows, start=2):
        for c in range(1, 12):
            md.cell(i, c)._style = copy(styles[c-1])
            md.cell(i, c).number_format = formats[c-1]
        vals = [
            r.get("month"), _excel_date(r.get("period_start")), _excel_date(r.get("period_end")), _excel_date(r.get("report_published")),
            r.get("platform"), r.get("policy_category"), r.get("content_actioned_raw"),
            r.get("content_actioned_numeric"), r.get("proactive_rate"), _proactive(r), r.get("total_user_grievances"),
        ]
        for c, v in enumerate(vals, 1):
            md.cell(i, c).value = v
        for c in (2, 3, 4):
            md.cell(i, c).number_format = "yyyy-mm-dd"
        for c in (8, 10, 11):
            md.cell(i, c).number_format = "#,##0"
        md.cell(i, 6).alignment = Alignment(vertical="center", wrap_text=True)
        md.row_dimensions[i].height = 45 if len(r.get("policy_category") or "") > 80 else 30
    for column in ("B", "C", "D"):
        md.column_dimensions[column].width = 14
    note_row = len(rows) + 3
    md.cell(note_row,1).value = "Note: 'Total User Grievances Received' is the platform-month total (not broken down by category) and is repeated on every category row for that platform-month, so it can be filtered/pivoted alongside category-level data. 'Proactive Actions by AI (est.)' = Content Actioned x Proactive Rate."
    md.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=11)
    md.freeze_panes = "A2"
    md.auto_filter.ref = f"A1:K{max(2, len(rows)+1)}"

    # Precompute monthly summary values. This avoids relying on Excel recalculation.
    groups = defaultdict(list)
    for r in rows:
        groups[(r.get("month"), r.get("period_end"), r.get("platform"))].append(r)
    summary = []
    for key, rs in sorted(groups.items(), key=lambda kv: (kv[0][1] or "", kv[0][2] or "")):
        total = int(round(sum(float(r.get("content_actioned_numeric") or 0) for r in rs)))
        proactive = sum(_proactive(r) for r in rs)
        rate = proactive / total if total else None
        grievances = next((r.get("total_user_grievances") for r in rs if r.get("total_user_grievances") is not None), None)
        summary.append((*key, total, proactive, rate, grievances))

    ms = wb["Monthly Summary"]
    summary_styles = [copy(ms.cell(2, c)._style) for c in range(1, 8)]
    if ms.max_row > 1:
        ms.delete_rows(2, ms.max_row - 1)
    for i, rec in enumerate(summary, start=2):
        for c, v in enumerate(rec, 1):
            ms.cell(i,c)._style = copy(summary_styles[c - 1])
            ms.cell(i,c).value = v
        ms.cell(i, 2).value = _excel_date(rec[1])
        ms.cell(i, 2).number_format = "yyyy-mm-dd"
    ms.freeze_panes = "A2"

    # Trend table: one row per month. Keep the template's two platform-specific charts.
    tc = wb["Trend Charts"]
    trend_styles = [copy(tc.cell(2, c)._style) for c in range(1, 7)]
    tc._charts = []
    if tc.max_row > 1:
        tc.delete_rows(2, tc.max_row - 1)
    month_map = defaultdict(dict)
    for month, period_end, platform, total, proactive, rate, grievances in summary:
        month_map[(month, period_end)][platform] = (total, grievances)
    for i, ((month, period_end), pmap) in enumerate(sorted(month_map.items(), key=lambda kv: kv[0][1] or ""), start=2):
        for c in range(1, 7):
            tc.cell(i,c)._style = copy(trend_styles[c - 1])
        tc.cell(i,1).value = month
        tc.cell(i,2).value = _excel_date(period_end)
        tc.cell(i,2).number_format = "yyyy-mm-dd"
        tc.cell(i,3).value = pmap.get("Facebook", (None,None))[0]
        tc.cell(i,4).value = pmap.get("Facebook", (None,None))[1]
        tc.cell(i,5).value = pmap.get("Instagram", (None,None))[0]
        tc.cell(i,6).value = pmap.get("Instagram", (None,None))[1]

    if month_map:
        end = len(month_map) + 1
        cats = Reference(tc, min_col=1, min_row=2, max_row=end)
        fb = LineChart(); fb.title = "Facebook: Content Actioned (AI) vs Grievances Received"; fb.y_axis.title="Count"; fb.x_axis.title="Month"; fb.height=10; fb.width=20
        fb.add_data(Reference(tc, min_col=3, max_col=4, min_row=1, max_row=end), titles_from_data=True)
        if len(fb.series) >= 2:
            fb.series[0].tx = SeriesLabel(v=tc.cell(1,3).value)
            fb.series[1].tx = SeriesLabel(v=tc.cell(1,4).value)
        fb.set_categories(cats); tc.add_chart(fb, "H1")
        ig = LineChart(); ig.title = "Instagram: Content Actioned (AI) vs Grievances Received"; ig.y_axis.title="Count"; ig.x_axis.title="Month"; ig.height=10; ig.width=20
        ig.add_data(Reference(tc, min_col=5, max_col=6, min_row=1, max_row=end), titles_from_data=True)
        if len(ig.series) >= 2:
            ig.series[0].tx = SeriesLabel(v=tc.cell(1,5).value)
            ig.series[1].tx = SeriesLabel(v=tc.cell(1,6).value)
        ig.set_categories(cats); tc.add_chart(ig, "H26")

    notes = wb["Data Notes"]
    for row_number, height in {1:45, 4:60, 7:45, 14:75, 15:60, 16:60, 18:60, 19:90, 22:90, 25:100}.items():
        notes.row_dimensions[row_number].height = height
    periods = [r.get("period_end") for r in rows if r.get("period_end")]
    platforms = sorted({r.get("platform") for r in rows if r.get("platform")})
    notes["A4"] = "Meta India Monthly Reports under the Information Technology (Intermediary Guidelines and Digital Media Ethics Code) Rules, 2021. Source: official Meta report PDFs discovered from Meta's Regulatory Transparency Reports hub."
    notes["A7"] = "Only reports that pass strict structural validation are included in Master Data. Suspicious/unsupported reports are written to the review folder instead of being silently accepted."
    notes["A8"] = f"{len(failed_reports)} report(s) have unresolved download or validation failures."
    notes["A25"] = "- Early 2021 reports use irregular reporting periods. Grievance totals are blank when the content-policy report does not publish them; totals from a separate reporting period are not substituted.\n- 'Proactive Actions by AI (est.)' is a derived estimate.\n- Numeric counts reflect Meta's published rounding; < and > values retain their threshold in the numeric column."
    notes["A9"] = "For counts reported as < or > a threshold, the raw column preserves the bound and the numeric column uses that threshold. Summary totals and proactive-action estimates inherit this approximation."
    notes.merge_cells("A9:H9")
    notes["A9"].alignment = Alignment(wrap_text=True, vertical="center")
    notes.row_dimensions[9].height = 45
    if periods:
        notes["A11"] = f"{len(set(r.get('month') for r in rows))} monthly data point(s), {min(periods)} – {max(periods)}. Platforms present: {', '.join(platforms)}."
    caveat_periods = sorted({r.get("month") for r in rows if r.get("source_notes")})
    if caveat_periods:
        notes.merge_cells("A12:H12")
        notes["A12"] = f"Source caveat for {', '.join(caveat_periods)}: Meta reports that grievance totals may be larger than stated because of a logging issue."
        notes["A12"].alignment = Alignment(wrap_text=True, vertical="center")
        notes.row_dimensions[12].height = 45
    for rr in range(28, 36):
        notes.cell(rr,1).value = None
        notes.cell(rr,3).value = None
    if not failed_reports:
        notes["A28"] = "None in the latest run."
    for j, item in enumerate(failed_reports[:8], start=28):
        text = _failed_report_to_text(item)
        label, separator, reason = text.partition(": ")
        notes.cell(j,1).value = label if separator else "Report"
        notes.cell(j,3).value = reason if separator else text
        notes.cell(j,3).alignment = Alignment(wrap_text=True, vertical="center")
        notes.row_dimensions[j].height = 45
    notes["A36"] = f"Generated by Meta India low-request pipeline on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}."

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
