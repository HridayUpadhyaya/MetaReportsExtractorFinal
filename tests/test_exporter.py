import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from src.exporter import build_workbook


class WorkbookExportTests(unittest.TestCase):
    def test_expansion_past_template_note_preserves_every_data_row(self):
        rows = [{
            "month": "Apr-2025", "period_start": "2025-04-01",
            "period_end": "2025-04-30", "report_published": "2025-05-31",
            "platform": "Facebook", "policy_category": f"Policy {index}",
            "policy_order": index, "content_actioned_raw": "1000",
            "content_actioned_numeric": 1000, "proactive_rate": 0.9,
            "total_user_grievances": 54320,
        } for index in range(60)]
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "report.xlsx"
            build_workbook(Path(__file__).resolve().parents[1] / "template.xlsx", output, rows)
            sheet = load_workbook(output)["Master Data"]
            exported = [row[5] for row in sheet.iter_rows(min_row=2, values_only=True) if row[4]]
            self.assertEqual(exported, [row["policy_category"] for row in rows])

    def test_summary_totals_dates_and_formats(self):
        rows = [
            {
                "month": "Apr-2025", "period_start": "2025-04-01",
                "period_end": "2025-04-30", "report_published": "2025-05-31",
                "platform": "Facebook", "policy_category": category,
                "policy_order": index, "content_actioned_raw": str(count),
                "content_actioned_numeric": count, "proactive_rate": rate,
                "total_user_grievances": 54320,
            }
            for index, (category, count, rate) in enumerate([
                ("Hate Speech", 1000, 0.90), ("Spam", 3000, 0.98)
            ], 1)
        ]
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "report.xlsx"
            build_workbook(Path(__file__).resolve().parents[1] / "template.xlsx", output, rows)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ["Data Notes", "Monthly Summary", "Trend Charts", "Master Data"])
            summary = workbook["Monthly Summary"]
            self.assertEqual(summary["D2"].value, 4000)
            self.assertEqual(summary["E2"].value, 3840)
            self.assertAlmostEqual(summary["F2"].value, 0.96)
            self.assertEqual(summary["G2"].value, 54320)
            self.assertEqual(summary["F2"].number_format, "0.0%")
            self.assertEqual(summary["D2"].number_format, "#,##0")
            self.assertEqual(workbook["Master Data"]["C2"].value.date().isoformat(), "2025-04-30")
            self.assertEqual(len(workbook["Trend Charts"]._charts), 2)


if __name__ == "__main__":
    unittest.main()
